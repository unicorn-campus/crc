# python check_paste.py {덱}.pptx — 본 덱에 붙여 넣기 전 점검(pptx-guide 6-13)
import sys
from pptx import Presentation
p = Presentation(sys.argv[1])
err = []
if (p.slide_width, p.slide_height) != (14630400, 8229600):
    err.append("슬라이드 크기가 16x9인치가 아님")
for i, s in enumerate(p.slides, 1):
    if s.slide_layout.name != "내용_단쪽":
        err.append(f"{i}쪽 레이아웃 이름: {s.slide_layout.name}")
    if s.shapes.title is None or not s.shapes.title.text.strip():
        err.append(f"{i}쪽 제목 개체 틀 비었음")
    for sh in s.shapes:
        if sh.has_text_frame and " / " in sh.text_frame.text and sh.top > 7_500_000 and len(sh.text_frame.text) < 10:
            err.append(f"{i}쪽 글자 쪽 표시 의심: {sh.text_frame.text}")
xml = "".join(s._element.xml for s in p.slides)
if "schemeClr" in xml:
    err.append("테마 색(schemeClr) 사용")
print("\n".join(err) or "OK")
